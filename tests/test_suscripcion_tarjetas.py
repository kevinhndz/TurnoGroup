from datetime import timedelta
from io import BytesIO
import pytest
from openpyxl import load_workbook
from backend.Suscripciones.tabla import TarjetasSuscripcion, ComprobantesSuscripcion, Suscripciones
from backend.Usuarios.tabla import Usuarios
from backend.Sesiones.tabla import Sesiones
from core.config import revision
from core.operaciones import ahora
from tests.test_roles_empresas import usar_rol
from utils.hash import encriptar_contrasena


def tarjeta(**extra):
    return {"tarjeta_prueba": "visa", "titular": "Persona de prueba", "mes": 12, "anio": 2099, **extra}


def test_tarjetas_prioridades_y_retiro(client, db, datos):
    primera = client.post('/suscripciones/tarjetas', json=tarjeta())
    assert primera.status_code == 201, primera.text
    assert primera.json()['prioridad'] == 'Principal'
    assert primera.json()['ultimos_cuatro'] == '4242'
    assert client.post('/suscripciones/tarjetas', json=tarjeta()).status_code == 409
    segunda = client.post('/suscripciones/tarjetas', json=tarjeta(tarjeta_prueba='mastercard')).json()
    assert client.post(f'/suscripciones/tarjetas/{segunda["id"]}/secundaria').status_code == 200
    assert client.delete(f'/suscripciones/tarjetas/{primera.json()["id"]}').status_code == 204
    actual = client.get('/suscripciones/actual').json()
    assert len(actual['tarjetas']) == 1
    assert actual['tarjetas'][0]['prioridad'] == 'Principal'
    assert client.delete(f'/suscripciones/tarjetas/{segunda["id"]}').status_code == 409
    assert client.patch('/suscripciones/renovacion', json={'renovacion_automatica':False}).status_code == 200
    assert client.delete(f'/suscripciones/tarjetas/{segunda["id"]}').status_code == 204
    assert client.patch('/suscripciones/renovacion', json={'renovacion_automatica':True}).status_code == 409
    assert db.query(TarjetasSuscripcion).filter_by(id_empresa=datos['empresa']).count() == 2


@pytest.mark.parametrize('extra', [{'numero':'4242424242424242'}, {'cvv':'123'}, {'mes':13}, {'anio':2026, 'mes':1}, {'tarjeta_prueba':'otra'}, {'titular':'  '}])
def test_tarjeta_invalida_no_guarda_datos(client, db, extra):
    assert client.post('/suscripciones/tarjetas', json=tarjeta(**extra)).status_code == 422
    assert db.query(TarjetasSuscripcion).count() == 0


def test_no_modifica_tarjeta_ni_descarga_comprobante_ajeno(client, db, datos):
    ajena = TarjetasSuscripcion(id_empresa=datos['otra_empresa'], titular='Otro titular', marca='Visa', ultimos_cuatro='4242', mes=12, anio=2099, prioridad='Principal')
    doc = ComprobantesSuscripcion(id_empresa=datos['otra_empresa'], nombre_empresa='Otra empresa', concepto='Prueba ajena', periodo_inicio=ahora(), periodo_fin=ahora(), importe=0, detalle='Prueba')
    db.add_all([ajena, doc]); db.commit()
    for accion in ['principal','secundaria']:
        assert client.post(f'/suscripciones/tarjetas/{ajena.id}/{accion}').status_code == 404
    assert client.delete(f'/suscripciones/tarjetas/{ajena.id}').status_code == 404
    assert client.get(f'/suscripciones/comprobantes/{doc.id}/pdf').status_code == 404
    assert client.get('/suscripciones/actual').json()['tarjetas'] == []


@pytest.mark.parametrize('rol', ['Mesero','Administrador','Superadministrador'])
def test_facturacion_exclusiva_del_gerente(client, db, datos, rol):
    usar_rol(client, db, datos, rol)
    assert client.get('/suscripciones/actual').status_code == 403
    assert client.post('/suscripciones/tarjetas', json=tarjeta()).status_code == 403
    assert client.get('/suscripciones/comprobantes/excel').status_code == 403
    assert client.post('/suscripciones/desactivar-cuenta', json={'clave':'x','confirmacion':'DESACTIVAR'}).status_code == 403


def test_registro_no_factura_cobros_inexistentes_y_exporta(client, db, datos):
    creado = client.post('/suscripciones/comprobantes/registro')
    assert creado.status_code == 201, creado.text
    id = creado.json()['id']
    assert float(creado.json()['importe']) == 0
    assert client.post('/suscripciones/comprobantes/registro').json()['id'] == id
    assert db.query(ComprobantesSuscripcion).count() == 1
    pdf = client.get(f'/suscripciones/comprobantes/{id}/pdf')
    assert pdf.status_code == 200 and pdf.content.startswith(b'%PDF-')
    assert 'application/pdf' in pdf.headers['content-type']
    xlsx = client.get('/suscripciones/comprobantes/excel')
    assert xlsx.status_code == 200
    hoja = load_workbook(BytesIO(xlsx.content)).active
    assert hoja.cell(3,7).value == 0
    assert 'SIN VALIDEZ FISCAL' in hoja['A1'].value


def preparar_clave(db, datos):
    usuario = db.get(Usuarios, datos['usuario_autorizado'])
    usuario.clave_protegida = encriptar_contrasena('Clave de prueba para desactivar 123')
    db.commit()
    return {'clave':'Clave de prueba para desactivar 123','confirmacion':'DESACTIVAR'}


def test_prueba_gratuita_no_permite_desactivar(client, db, datos):
    cuerpo = preparar_clave(db, datos)
    assert client.post('/suscripciones/comprobantes/registro').status_code == 201
    respuesta = client.post('/suscripciones/desactivar-cuenta', json=cuerpo)
    assert respuesta.status_code == 409
    assert db.get(Usuarios, datos['usuario_autorizado']).estado
    assert db.get(Sesiones, datos['sesion']).fecha_revocacion is None


def test_cuenta_pagada_se_desactiva_y_reactiva_con_confirmacion(client, db, datos):
    cuerpo = preparar_clave(db, datos)
    db.add(ComprobantesSuscripcion(id_empresa=datos['empresa'], nombre_empresa='Empresa', concepto='Mensualidad de prueba', periodo_inicio=ahora()-timedelta(days=30), periodo_fin=ahora(), importe=100, detalle='Pago simulado para comprobar la desactivacion'))
    db.commit()
    assert client.post('/suscripciones/desactivar-cuenta', json={**cuerpo,'clave':'incorrecta'}).status_code == 403
    assert client.post('/suscripciones/desactivar-cuenta', json=cuerpo).status_code == 204
    assert not db.get(Usuarios,datos['usuario_autorizado']).estado
    assert not db.query(Suscripciones).filter_by(id_empresa=datos['empresa']).one().renovacion_automatica
    assert client.get('/sesiones/actual').status_code == 401
    acceso = {'correo':'gerente@example.com','clave':cuerpo['clave']}
    assert client.post('/sesiones/iniciar', json={**acceso,'clave':'incorrecta'}).status_code == 401
    respuesta = client.post('/sesiones/iniciar', json=acceso)
    assert respuesta.status_code == 409
    assert respuesta.json()['detail']['codigo'] == 'CUENTA_DESACTIVADA'
    assert client.post('/sesiones/iniciar', json={**acceso,'reactivar':True}).status_code == 200
    assert client.get('/sesiones/actual').status_code == 200
    assert db.get(Usuarios,datos['usuario_autorizado']).fecha_desactivacion_voluntaria is None


def test_bloqueo_administrativo_no_se_reactiva(client, db, datos):
    cuerpo = preparar_clave(db, datos)
    usuario = db.get(Usuarios, datos['usuario_autorizado']); usuario.estado = False; db.commit()
    assert client.post('/sesiones/iniciar', json={'correo':usuario.correo,'clave':cuerpo['clave'],'reactivar':True}).status_code == 401


def test_modo_paddle_no_acepta_simulaciones(client, monkeypatch):
    monkeypatch.setattr(revision, 'FACTURACION_MODO', 'paddle')
    assert client.post('/suscripciones/tarjetas', json=tarjeta()).status_code == 503


@pytest.mark.parametrize('estado',['Prueba','Activa'])
def test_cancelacion_conserva_acceso_hasta_el_fin_del_periodo(client, db, datos, estado):
    suscripcion = db.query(Suscripciones).filter_by(id_empresa=datos['empresa']).one()
    suscripcion.estado = estado
    fin = ahora() + timedelta(days=3)
    if estado == 'Prueba':
        suscripcion.fin_prueba = fin
    else:
        suscripcion.proximo_cobro = fin
    db.commit()
    respuesta = client.patch('/suscripciones/renovacion', json={'renovacion_automatica':False})
    assert respuesta.status_code == 200, respuesta.text
    assert respuesta.json()['estado'] == estado
    assert db.get(Usuarios, datos['usuario_autorizado']).estado
    assert client.get('/pedidos/').status_code == 200
    assert suscripcion.fin_prueba == fin if estado == 'Prueba' else suscripcion.proximo_cobro == fin
    if estado == 'Prueba':
        suscripcion.fin_prueba = ahora() - timedelta(seconds=1)
    else:
        suscripcion.proximo_cobro = ahora() - timedelta(seconds=1)
    db.commit()
    assert client.get('/pedidos/').status_code == 403
    assert client.get('/sesiones/actual').status_code == 200


def test_administrador_puede_impedir_reactivacion_de_cuenta_voluntaria(client, db, datos):
    cuerpo = preparar_clave(db, datos)
    usuario = db.get(Usuarios, datos['usuario_autorizado'])
    usuario.fecha_desactivacion_voluntaria = ahora()
    usuario.estado = False
    db.commit()
    from backend.Usuarios.service import UsuarioService
    from backend.Usuarios.schema import Revisar_Json_Editar_Usuario
    UsuarioService.editar_usuario_service(Revisar_Json_Editar_Usuario(estado=False), usuario.id, db)
    assert usuario.fecha_desactivacion_voluntaria is None
    assert client.post('/sesiones/iniciar',json={'correo':usuario.correo,'clave':cuerpo['clave'],'reactivar':True}).status_code == 401
