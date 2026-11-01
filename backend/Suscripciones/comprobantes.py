from io import BytesIO
from textwrap import wrap


def crear_pdf(item):
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A4

    salida = BytesIO()
    pdf = canvas.Canvas(salida, pagesize=A4)
    pdf.setTitle(f"Comprobante de prueba TG-{item.id}")
    pdf.setFillColorRGB(.60, .25, .14)
    pdf.setFont("Helvetica-Bold", 23)
    pdf.drawString(50, 785, "Turno Group")
    pdf.setFillColorRGB(.15, .15, .15)
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(50, 745, "COMPROBANTE DE PRUEBA - SIN VALIDEZ FISCAL")
    lineas = [
        f"Numero: TG-PRUEBA-{item.id:06}",
        f"Empresa: {item.nombre_empresa}",
        f"Fecha: {item.fecha:%d/%m/%Y}",
        f"Concepto: {item.concepto}",
        f"Periodo: {item.periodo_inicio:%d/%m/%Y} al {item.periodo_fin:%d/%m/%Y}",
        "",
        item.detalle,
        "",
        f"Total: {item.importe:.2f} {item.moneda}",
        "No acredita pagos reales. Documento generado para probar la interfaz.",
    ]
    pdf.setFont("Helvetica", 11)
    y = 707
    for linea in lineas:
        for parte in wrap(linea, width=80) or [""]:
            pdf.drawString(50, y, parte)
            y -= 21
    pdf.save()
    return salida.getvalue()


def crear_excel(filas):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Comprobantes de prueba"
    hoja.append(["DOCUMENTOS DE PRUEBA - SIN VALIDEZ FISCAL"])
    hoja.append(["Numero", "Empresa", "Fecha", "Concepto", "Inicio", "Fin", "Importe", "Moneda", "Detalle"])
    for item in filas:
        hoja.append([f"TG-PRUEBA-{item.id:06}", item.nombre_empresa, item.fecha, item.concepto, item.periodo_inicio, item.periodo_fin, float(item.importe), item.moneda, item.detalle])
        for celda in hoja[hoja.max_row]:
            if isinstance(celda.value, str):
                celda.data_type = "s"
        for columna in (3, 5, 6):
            hoja.cell(hoja.max_row, columna).number_format = "dd/mm/yyyy"
    for celda in hoja[2]:
        celda.font = Font(bold=True)
    hoja.freeze_panes = "A3"
    for columna in "ABCDEFGHI":
        hoja.column_dimensions[columna].width = 24 if columna != "I" else 65
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()
