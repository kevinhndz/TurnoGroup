from fastapi import HTTPException
from sqlalchemy.orm import Session
from backend.Categorias.repository import CategoriaRepository as repo_categoria
from backend.Empresas.repository import EmpresaRepository as repo_empresa
from backend.Productos.repository import ProductoRepository as repo
from backend.Productos.schema import (
    Revisar_Json_Nuevo_Producto,
    Revisar_Json_Editar_Producto,
)
from backend.Productos.tabla import Productos
from core.exceptions import RecursoNoEncontradoError, RecursoDuplicadoError


class ProductoService:

    @staticmethod
    def revisar_archivo_service(json):
        import base64
        import binascii
        import csv
        from io import BytesIO, StringIO
        from zipfile import ZipFile
        from decimal import Decimal, InvalidOperation

        extension = ".xlsx" if json.formato == "excel" else ".csv"
        if not json.nombre_archivo.lower().endswith(extension):
            raise HTTPException(422, f"Selecciona un archivo {extension}")
        try:
            contenido = base64.b64decode(json.contenido, validate=True)
        except (ValueError, binascii.Error):
            raise HTTPException(422, "No se pudo leer el archivo") from None
        if len(contenido) > 1_000_000:
            raise HTTPException(413, "El archivo supera 1 MB")
        if json.formato == "excel":
            from openpyxl import load_workbook
            libro = None
            try:
                with ZipFile(BytesIO(contenido)) as archivo:
                    entradas = archivo.infolist()
                    if len(entradas) > 1000 or sum(item.file_size for item in entradas) > 10_000_000:
                        raise HTTPException(422, "El Excel contiene demasiada informacion")
                    for item in entradas:
                        if item.filename.lower().endswith(".xml"):
                            xml = archivo.read(item).upper()
                            if b"<!DOCTYPE" in xml or b"<!ENTITY" in xml:
                                raise HTTPException(422, "Usa un Excel sin contenido XML personalizado")
                libro = load_workbook(BytesIO(contenido), read_only=True, data_only=False, keep_links=False)
                hoja = libro.worksheets[0]
                if hoja.max_row and hoja.max_row > 1001:
                    raise HTTPException(422, "El Excel debe tener hasta 1000 productos y una fila de encabezados")
                filas = []
                for fila in hoja.iter_rows(max_row=1002, max_col=21):
                    if any(celda.data_type == "f" for celda in fila):
                        raise HTTPException(422, "Reemplaza las formulas del Excel por sus valores antes de importar")
                    valores = [celda.value for celda in fila]
                    while valores and valores[-1] is None:
                        valores.pop()
                    if any(valor is not None and str(valor).strip() for valor in valores):
                        filas.append(valores)
                if hoja.max_column and hoja.max_column > 20:
                    raise HTTPException(422, "Usa solo las columnas de la plantilla")
            except HTTPException:
                raise
            except Exception:
                raise HTTPException(422, "No se pudo abrir el Excel. Usa un archivo .xlsx valido, sin contrasena") from None
            finally:
                if libro:
                    libro.close()
        else:
            try:
                texto = contenido.decode("utf-8-sig")
                try:
                    separador = csv.Sniffer().sniff(texto[:4096], delimiters=",;").delimiter
                except csv.Error:
                    separador = ","
                filas = []
                for fila in csv.reader(StringIO(texto), delimiter=separador, strict=True):
                    if any(valor.strip() for valor in fila):
                        filas.append(fila)
                    if len(filas) > 1001:
                        break
            except (UnicodeDecodeError, csv.Error):
                raise HTTPException(422, "No se pudo leer el CSV. Guardalo como CSV UTF-8") from None
        if not 2 <= len(filas) <= 1001:
            raise HTTPException(422, "El archivo debe tener entre 1 y 1000 productos")
        titulos = [str(valor or "").strip().lower() for valor in filas[0]]
        permitidos = {"nombre", "precio", "categoria", "descripcion"}
        if not {"nombre", "precio"}.issubset(titulos) or len(set(titulos)) != len(titulos) or not set(titulos).issubset(permitidos):
            raise HTTPException(422, "Usa las columnas nombre, precio, categoria y descripcion, sin repetir. Nombre y precio son obligatorias")
        productos = []
        nombres = set()
        for numero, valores in enumerate(filas[1:], start=2):
            if len(valores) > len(titulos):
                raise HTTPException(422, f"Revisa las columnas de la fila {numero}")
            valores = valores + [None] * (len(titulos) - len(valores))
            fila = {titulo: str(valor).strip() if valor is not None else "" for titulo, valor in zip(titulos, valores)}
            try:
                precio = Decimal(fila["precio"])
                if not precio.is_finite() or precio < 0 or precio > Decimal("99999999.99") or precio != precio.quantize(Decimal("0.01")):
                    raise ValueError()
            except (InvalidOperation, ValueError):
                raise HTTPException(422, f"Revisa el precio de la fila {numero}. Usa un numero no negativo con hasta dos decimales") from None
            if not fila["nombre"] or len(fila["nombre"]) > 100 or len(fila.get("descripcion", "")) > 10000:
                raise HTTPException(422, f"Revisa el nombre y la descripcion de la fila {numero}")
            if fila["nombre"].casefold() in nombres:
                raise HTTPException(409, f"El producto de la fila {numero} esta repetido en el archivo")
            nombres.add(fila["nombre"].casefold())
            productos.append({"nombre": fila["nombre"], "precio": format(precio, ".2f"), "categoria": fila.get("categoria", ""), "descripcion": fila.get("descripcion") or None})
        return {"productos": productos}

    @staticmethod
    def crear_producto_service(
        json: Revisar_Json_Nuevo_Producto, db: Session, id_empresa: int
    ):
        empresa = repo_empresa.buscar_empresa_por_id(db, id_empresa)
        if empresa is None:
            raise RecursoNoEncontradoError("No se encontro la empresa")

        if json.id_categoria is not None:
            categoria = repo_categoria.buscar_categoria_por_id(
                db, json.id_categoria, id_empresa
            )
            if categoria is None or not categoria.activo:
                raise RecursoNoEncontradoError(
                    "No se encontro la categoria de esta empresa"
                )

        look_up = repo.buscar_producto_por_nombre(db, id_empresa, json.nombre)
        if look_up is not None:
            raise RecursoDuplicadoError(
                "Ya existe un producto con ese nombre en la empresa"
            )

        nuevo_producto = Productos(
            id_empresa=id_empresa,
            id_categoria=json.id_categoria,
            nombre=json.nombre,
            descripcion=json.descripcion,
            precio=json.precio,
        )
        return repo.guardar_nuevo_producto(db, nuevo_producto)

    @staticmethod
    def listar_productos_service(
        id_empresa: int, pagina_actual: int, limite: int, db: Session, id_sucursal=None, filtros=None):
        empresa = repo_empresa.buscar_empresa_por_id(db, id_empresa)
        if empresa is None:
            raise RecursoNoEncontradoError("No se encontro la empresa")

        total, lista_productos = repo.listar_productos(
            db, id_empresa, pagina_actual, limite, id_sucursal, filtros=filtros)
        total_paginas = (total + limite - 1) // limite
        return {
            "total": total,
            "pagina_actual": pagina_actual,
            "limite": limite,
            "total_paginas": total_paginas,
            "data": lista_productos,
        }

    @staticmethod
    def ver_producto_service(id: int, id_empresa: int, db: Session, id_sucursal=None):
        empresa = repo_empresa.buscar_empresa_por_id(db, id_empresa)
        if empresa is None:
            raise RecursoNoEncontradoError("No se encontro la empresa")

        look_up = repo.buscar_producto_por_id(db, id, id_empresa, id_sucursal)

        if look_up is None:
            raise RecursoNoEncontradoError("No se encontro el producto")

        return look_up

    @staticmethod
    def editar_producto_service(
        json: Revisar_Json_Editar_Producto, id: int, id_empresa: int, db: Session
    ):
        empresa = repo_empresa.buscar_empresa_por_id(db, id_empresa)
        if empresa is None:
            raise RecursoNoEncontradoError("No se encontro la empresa")

        look_up = repo.buscar_producto_por_id(db, id, id_empresa)

        if look_up is None:
            raise RecursoNoEncontradoError("No se encontro el producto")

        cambios = json.model_dump(exclude_unset=True)
        if not cambios:
            raise HTTPException(
                status_code=422, detail="Debes enviar al menos un campo para editar"
            )
        for campo, valor in cambios.items():
            if valor is None and campo not in ("id_categoria", "descripcion"):
                raise HTTPException(
                    status_code=422, detail="El campo " + campo + " no acepta null"
                )

        if "nombre" in cambios and cambios["nombre"] != look_up.nombre:
            producto_existente = repo.buscar_producto_por_nombre(
                db, id_empresa, cambios["nombre"]
            )
            if producto_existente is not None:
                raise RecursoDuplicadoError(
                    "Ya existe un producto con ese nombre en la empresa"
                )

        if "id_categoria" in cambios and cambios["id_categoria"] is not None:
            categoria = repo_categoria.buscar_categoria_por_id(
                db, cambios["id_categoria"], id_empresa
            )
            if categoria is None or not categoria.activo:
                raise RecursoNoEncontradoError(
                    "No se encontro la categoria de esta empresa"
                )

        # Recorre el diccionario y actualiza cada campo enviado.
        for campo, valor in cambios.items():
            if valor is not None:
                setattr(look_up, campo, valor)
            elif campo == "id_categoria" or campo == "descripcion":
                setattr(look_up, campo, None)

        db.commit()
        db.refresh(look_up)
        return look_up

    @staticmethod
    def desactivar_producto_service(id: int, id_empresa: int, db: Session):
        empresa = repo_empresa.buscar_empresa_por_id(db, id_empresa)
        if empresa is None:
            raise RecursoNoEncontradoError("No se encontro la empresa")

        look_up = repo.buscar_producto_por_id(db, id, id_empresa)

        if look_up is None:
            raise RecursoNoEncontradoError("No se encontro el producto")

        look_up.disponible = False
        db.commit()
        db.refresh(look_up)
        return look_up

    @staticmethod
    def importar_productos_service(json, db, id_empresa, id_sucursal):
        from backend.DisponibilidadProductos.tabla import DisponibilidadProductos
        from backend.Sucursales.tabla import Sucursales
        from sqlalchemy import func

        nombres = [item.nombre.casefold() for item in json.productos]
        if len(set(nombres)) != len(nombres):
            raise HTTPException(409, "Hay productos repetidos en el archivo")
        existentes = {
            nombre.casefold()
            for nombre, in db.query(Productos.nombre)
            .filter(Productos.id_empresa == id_empresa)
            .all()
        }
        if existentes.intersection(nombres):
            raise HTTPException(409, "Un producto del archivo ya existe en tu catalogo")
        for item in json.productos:
            if item.id_categoria:
                categoria = repo_categoria.buscar_categoria_por_id(
                    db, item.id_categoria, id_empresa
                )
                if not categoria or not categoria.activo:
                    raise HTTPException(
                        404, "No se encontro la categoria de esta empresa"
                    )
        if json.asignar_sucursal:
            sucursal = (
                db.query(Sucursales)
                .filter(
                    Sucursales.id == id_sucursal,
                    Sucursales.id_empresa == id_empresa,
                    Sucursales.activo.is_(True),
                )
                .first()
            )
            if not sucursal:
                raise HTTPException(404, "No se encontro la sucursal")
        productos = [
            Productos(id_empresa=id_empresa, **item.model_dump())
            for item in json.productos
        ]
        db.add_all(productos)
        db.flush()
        if json.asignar_sucursal:
            db.add_all(
                [
                    DisponibilidadProductos(
                        id_producto=item.id, id_sucursal=id_sucursal, disponible=True
                    )
                    for item in productos
                ]
            )
        respuesta = {"creados": len(productos)}
        db.commit()
        return respuesta

    @staticmethod
    def guardar_imagen_service(json, id, id_empresa, db):
        import base64
        import binascii
        import io
        import warnings
        from pathlib import Path
        from uuid import uuid4
        from PIL import Image, ImageOps, UnidentifiedImageError

        look_up = repo.buscar_producto_por_id(db, id, id_empresa)
        if not look_up:
            raise RecursoNoEncontradoError("No se encontro el producto")
        try:
            contenido = base64.b64decode(json.contenido, validate=True)
            if len(contenido) > 5 * 1024 * 1024:
                raise HTTPException(413, "La imagen supera 5 MB")
            with warnings.catch_warnings():
                warnings.simplefilter("error", Image.DecompressionBombWarning)
                with Image.open(io.BytesIO(contenido)) as imagen:
                    if (
                        imagen.format not in ("JPEG", "PNG")
                        or imagen.width * imagen.height > 16_000_000
                    ):
                        raise HTTPException(
                            422, "Usa JPG o PNG de hasta 16 millones de pixeles"
                        )
                    imagen = ImageOps.exif_transpose(imagen).convert("RGB")
                    imagen.thumbnail((1600, 1600))
                    salida = io.BytesIO()
                    imagen.save(salida, format="JPEG", quality=85)
        except (
            binascii.Error,
            ValueError,
            UnidentifiedImageError,
            OSError,
            Image.DecompressionBombWarning,
            Image.DecompressionBombError,
        ):
            raise HTTPException(
                422, "El archivo no es una imagen JPG o PNG valida"
            ) from None
        carpeta = ProductoService.carpeta_imagenes(id_empresa)
        carpeta.mkdir(parents=True, exist_ok=True)
        nombre = str(uuid4()) + ".jpg"
        archivo = carpeta / nombre
        archivo.write_bytes(salida.getvalue())
        anterior = look_up.ruta_imagen
        try:
            look_up.ruta_imagen = nombre
            db.commit()
        except Exception:
            archivo.unlink(missing_ok=True)
            raise
        if anterior and Path(anterior).name == anterior:
            (carpeta / anterior).unlink(missing_ok=True)
        return look_up

    @staticmethod
    def carpeta_imagenes(id_empresa):
        import os
        from pathlib import Path

        return Path(
            os.environ.get(
                "TURNO_ARCHIVOS", str(Path(__file__).resolve().parents[2] / ".archivos")
            )
        ) / str(id_empresa)

    @staticmethod
    def consultar_imagen_service(id, id_empresa, db, id_sucursal):
        from pathlib import Path

        look_up = ProductoService.ver_producto_service(id, id_empresa, db, id_sucursal)
        if (
            not look_up.ruta_imagen
            or Path(look_up.ruta_imagen).name != look_up.ruta_imagen
        ):
            raise RecursoNoEncontradoError("No se encontro la imagen")
        archivo = ProductoService.carpeta_imagenes(id_empresa) / look_up.ruta_imagen
        if not archivo.is_file():
            raise RecursoNoEncontradoError("No se encontro la imagen")
        return archivo
